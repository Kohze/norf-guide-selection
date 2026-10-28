from pathlib import Path
import datetime,hashlib,json
import numpy as np
import pandas as pd
import openpyxl
M=Path(__file__).resolve().parent
N=M.parent if (M.parent/'cycle_87_fresh_guide_cross_screen_nORF_validation').exists() else M.parents[1]/'exploration_2026-10-11'/'novelty_discovery'
C87=N/'cycle_87_fresh_guide_cross_screen_nORF_validation'
S=M/'sources'
checks={}
def ck(name,value):
 checks[name]=bool(value)
 if not value:raise AssertionError(name)
def read(sheet):
 w=openpyxl.load_workbook(S/'41587_2020_806_MOESM5_ESM.xlsx',read_only=True,data_only=True)
 rows=w[sheet].iter_rows(min_row=3,values_only=True);headers=next(rows)
 return [dict(zip(headers,row)) for row in rows]
meta={r['Construct Barcode']:r for r in read('Table S27')}
lib=read('Table S24');assign={}
for r in lib:assign.setdefault(r['sgRNA Sequence'],set()).add(r['Input'])
w=openpyxl.load_workbook(S/'41587_2020_806_MOESM12_ESM.xlsx',read_only=True,data_only=True)
rows=w.active.iter_rows(values_only=True);headers=next(rows);raw=pd.DataFrame(rows,columns=headers)
rawmap=raw.set_index('Construct Barcode')
new=pd.read_csv(C87/'all720_newer_exact_product_guide_roles_and_eligibility_before_outcome_join.csv')
expected=[]
for _,r in new.iterrows():
 s=r.sgrna;t=meta.get(s)
 if not t or not r.primary_eligible or t['Gene_type']!='ORF_exon':continue
 import re
 name=re.sub(r'_exon[0-9]+$','',str(t['Gene Symbol']))
 if name!=r.old_source_name or s not in rawmap.index or rawmap.loc[s,'pDNA']<10 or len(assign.get(s,set()))!=1:continue
 expected.append((r.pair_ID,s))
md=pd.read_csv(M/'all720_followup_guide_metadata_eligibility.csv')
observed=md[md.secondary_metadata_eligible][['pair_ID','sgrna']]
ck('metadata_eligible_pairs_exactly_reconstructed',set(expected)==set(map(tuple,observed.values)))
n=pd.Series([p for p,g in expected]).value_counts()
exp_main={(p,g) for p,g in expected if n[p]>=3}
d=pd.read_csv(M/'primary_followup_paired_residuals.csv')
ck('primary_72_guide_memberships_exactly_reconstructed',exp_main==set(map(tuple,d[['pair_ID','sgrna']].values)))
ck('primary_16_products_72_guides_144_rows',(d.pair_ID.nunique(),d.sgrna.nunique(),len(d))==(16,72,144))
cells=['A375','HEPG2'];pub={'A375':'LFC A375 pDNA_D21','HEPG2':'LFC HepG2 pDNA_D21'}
source=np.array([meta[r.sgrna][pub[r.cell]] for _,r in d.iterrows()])
ck('all_primary_secondary_outcomes_match_original_workbook',np.max(np.abs(source-d.secondary_LFC))<1e-12)
latest=pd.read_csv(C87/'all2160_newer_exact_product_source_guide_cell_outcomes_and_roles.csv').set_index(['pair_ID','sgrna','cell'])
y=np.array([latest.loc[(r.pair_ID,r.sgrna,r.cell),'LFC'] for _,r in d.iterrows()])
ck('all_primary_newer_outcomes_match_pinned_source',np.max(np.abs(y-d.newer_LFC))<1e-12)
calc={}
for cell,prefix in [('A375','A375 '),('HEPG2','HepG2 ')]:
 end=[c for c in raw if c.startswith(prefix) and c.endswith('D21')]
 terms=[]
 for col in end:terms.append(np.log2(1+raw[col].to_numpy()/raw[col].sum()*1e6))
 initial=np.log2(1+raw.pDNA.to_numpy()/raw.pDNA.sum()*1e6)
 calc[cell]=dict(zip(raw['Construct Barcode'],np.mean(terms,axis=0)-initial))
direct=np.array([calc[r.cell][r.sgrna] for _,r in d.iterrows()])
ck('published_secondary_endpoint_reconstructed_from_original_counts',np.max(np.abs(direct-source))<2e-12)
x=source-d.assign(original=source).groupby(['pair_ID','cell']).original.transform('mean').to_numpy()
y=y-d.assign(original=y).groupby(['pair_ID','cell']).original.transform('mean').to_numpy()
w=1/(2*d.groupby('pair_ID').sgrna.transform('nunique').to_numpy())
def corr(a,b,weight):
 c=np.cov(a,b,aweights=weight,ddof=0)
 return float(c[0,1]/np.sqrt(c[0,0]*c[1,1]))
value=corr(x,y,w);result=json.loads((M/'FOLLOWUP_GUIDE_RESULTS.json').read_text())
ck('independent_weighted_covariance_matches_primary',abs(value-result['primary']['correlation'])<1e-12)
ck('residuals_and_weights_match',np.allclose(x,d.x,atol=1e-12) and np.allclose(y,d.y,atol=1e-12) and np.allclose(w,d.weight,atol=1e-12))
plans=np.load(M/'all2000_followup_cluster_draws.npz');clusters=plans['cluster_ids'];ix=[np.flatnonzero(d.dependence_cluster.to_numpy()==c) for c in clusters]
boot=[]
for draw in plans['draws']:
 idx=np.concatenate([ix[j] for j in draw]);boot.append(corr(x[idx],y[idx],w[idx]))
stored=pd.read_csv(M/'all2000_followup_cluster_bootstrap.csv').correlation.to_numpy()
ck('all2000_bootstrap_draws_replayed_by_explicit_row_duplication',np.max(np.abs(stored-boot))<1e-12)
ck('percentile_CI_reproduced',np.allclose(np.quantile(boot,[.025,.975]),result['primary']['nominal95_cluster_CI'],atol=1e-12))
rng=np.random.default_rng(20261017);mapping=[]
for pid in sorted(d.pair_ID.unique()):
 q=d[d.pair_ID.eq(pid)];guides=sorted(q.sgrna.unique())
 mapping.append(np.array([[q.index[(q.sgrna==g)&(q.cell==c)][0] for c in cells] for g in guides]))
perm=[]
for j in range(5000):
 yy=y.copy()
 for indexes in mapping:yy[indexes.ravel()]=y[indexes[rng.permutation(len(indexes))].ravel()]
 perm.append(corr(x,yy,w))
storedperm=pd.read_csv(M/'all5000_followup_identity_permutations.csv').correlation.to_numpy()
ck('all5000_permutations_replayed_by_explicit_row_mapping',np.max(np.abs(storedperm-perm))<1e-12)
ck('permutation_P_reproduced',(1+np.sum(np.abs(perm)>=abs(value)))/5001==result['primary']['permutation_two_sided_add_one_P'])
# Source-level count fingerprints: shared sequence identity is not sample identity.
b=pd.read_csv(N/'cycle_35_guide_replicates_and_medical_provenance'/'all_original_table13_B_guide_measurements.csv').set_index('sgrna')
shared=sorted(set(rawmap.index)&set(b.index));fingerprints=[]
for sec in ['pDNA']+[c for c in raw if c.startswith(('A375 ','HepG2 ')) and c.endswith('D21')]:
 for newer in [c for c in b if c.startswith(('pDNA','A375','HEPG2'))]:
  aa=rawmap.loc[shared,sec].to_numpy(dtype=float);bb=b.loc[shared,newer].to_numpy(dtype=float)
  fingerprints.append({'secondary_sample':sec,'newer_sample':newer,'shared_guides':len(shared),'identical_count_rows':int(np.sum(aa==bb)),'all_counts_identical':bool(np.array_equal(aa,bb)),'log_count_Pearson':float(np.corrcoef(np.log1p(aa),np.log1p(bb))[0,1])})
ck('no_secondary_newer_sample_has_identical_shared_count_vector',not any(r['all_counts_identical'] for r in fingerprints))
pd.DataFrame(fingerprints).to_csv(M/'secondary_newer_sample_count_fingerprints.csv',index=False)
pins=json.loads((M/'METADATA_PREFLIGHT.json').read_text())['metadata_pins']
ck('frozen_source_hashes_unchanged',all(hashlib.sha256(Path(p['path']).read_bytes()).hexdigest()==p['sha256'] for p in pins))
report={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'status':'pass','verification':'Separate computational implementation by the same assistant, not independent biological or expert review.','checks':checks,'passed_checks':sum(checks.values()),'primary_correlation':value,'max_bootstrap_difference':float(np.max(np.abs(stored-boot))),'max_permutation_difference':float(np.max(np.abs(storedperm-perm))),'raw_count_endpoint_max_absolute_difference':float(np.max(np.abs(direct-source))),'source_provenance':'2021 primary and follow-up assays are in the same publication;2026 Methods distinguish the current experiment from older reprocessed studies. No exact matching secondary/newer raw count sample vector was found. This fingerprint alone cannot prove biological independence.'}
(M/'SEPARATE_COMPUTATIONAL_VERIFICATION.json').write_text(json.dumps(report,indent=2),encoding='utf8')
print(json.dumps(report,indent=2))
