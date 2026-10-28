"""Inspect newly identified processed supplement schemas, not fitness values."""
from pathlib import Path
import datetime as dt
import hashlib
import json
import urllib.request
import urllib.error
import openpyxl

OUT = Path(__file__).resolve().parent
SRC = OUT / 'sources'

def dump(path, obj):
    path.write_text(json.dumps(obj, indent=2) + '\n', encoding='utf-8')

def main():
    amendment = OUT / 'INTAKE_ADDENDUM_INDIVIDUAL_GUIDE_SUPPLEMENTS.json'
    if not amendment.exists():
        dump(amendment, {
            'utc': dt.datetime.now(dt.timezone.utc).isoformat(),
            'reason': 'Primary indexed supplementary-data listing identifies TableS3 individual-sgRNA results and TableS4 library design, outside the previously inspected author repository.',
            'listing_source': 'https://pmc.ncbi.nlm.nih.gov/articles/PMC11929002/ (search index only; live page challenged; no retry)',
            'publisher_article_identity': 'https://www.sciencedirect.com/science/article/pii/S2589004225001440',
            'resource_purpose': 'Acquire the processed data supplements themselves, not any challenged article text/PDF.',
            'additional_bound_bytes': 30_000_000,
            'files': ['1-s2.0-S2589004225001440-mmc4.xlsx', '1-s2.0-S2589004225001440-mmc5.xlsx'],
            'stage': 'First-row schema and library identity only; no fitness value comparison.',
            'conditional_plan': 'If individual-guide A375 outcomes and exact20nt sequences are available, freeze a third-study within-product shared-guide concordance protocol before opening numeric outcomes.',
            'no_access_bypass': True,
            'not_independent_preregistration': True
        })
    total = 0
    reports = []
    for suffix in [4, 5]:
        name = f'1-s2.0-S2589004225001440-mmc{suffix}.xlsx'
        dest = SRC / name
        url = f'https://ars.els-cdn.com/content/image/{name}'
        rec = SRC / (name + '.response.json')
        if not dest.exists() and not rec.exists():
            try:
                with urllib.request.urlopen(url, timeout=40) as response:
                    data = response.read(30_000_001 - total)
                    info = {'status': response.status, 'headers': dict(response.headers)}
            except urllib.error.HTTPError as error:
                data = error.read(10000)
                info = {'status': error.code, 'headers': dict(error.headers), 'access_boundary_no_retry': True}
            info.update({'utc': dt.datetime.now(dt.timezone.utc).isoformat(), 'url': url,
                         'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()})
            dump(rec, info)
            if info['status'] != 200:
                (SRC / (name + '.unavailable.body')).write_bytes(data)
                print(json.dumps(info))
                continue
            if total + len(data) > 30_000_000:
                raise RuntimeError('Additional data bound exceeded')
            assert data.startswith(b'PK'), 'Not an XLSX ZIP'
            dest.write_bytes(data)
        if not dest.exists():
            reports.append({'file': name, 'status': 'unavailable_no_retry'})
            continue
        total += dest.stat().st_size
        wb = openpyxl.load_workbook(dest, read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            sheets.append({'sheet': ws.title, 'rows': ws.max_row, 'cols': ws.max_column, 'first_row_only': list(first)})
        reports.append({'file': name, 'sheets': sheets, 'outcomes_loaded': False})
        wb.close()
    dump(OUT / 'SUPPLEMENT_SCHEMA_ONLY.json', {'utc': dt.datetime.now(dt.timezone.utc).isoformat(), 'reports': reports})
    print(json.dumps(reports, indent=2))

if __name__ == '__main__':
    main()
